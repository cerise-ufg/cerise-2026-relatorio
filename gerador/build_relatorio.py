#!/usr/bin/env python3
"""Gera o relatório CERISE (.docx) a partir de um Markdown, usando o modelo oficial como base.

Mantém a capa, o cabeçalho com os logos, a numeração automática dos títulos (1., 1.1., 1.1.1.)
e o estilo de parágrafo do modelo (Arial 12, justificado, recuo de 1,25 cm na primeira linha).

Markdown suportado:
  # / ## / ###            títulos numerados (níveis 1-3)
  parágrafos              linhas consecutivas viram um parágrafo; linha em branco separa
  - item / * item         marcadores
  ![legenda](caminho){w=14}   figura centralizada (largura em cm, padrão 15) + "Figura N. legenda"
  Tabela: legenda         linha logo antes de uma tabela pipe -> "Tabela N. legenda" acima dela
  | a | b |               tabela pipe (1ª linha = cabeçalho; linha |---| ignorada)
  **negrito** *itálico* `código` [texto](url)
  <!-- pagebreak -->      quebra de página
  [1] referência          parágrafo de referência (sem recuo)

Paginação (revisão de 26/10/2026):
  - títulos alinhados à esquerda, com o número seguido de um espaço (sem tabulação);
  - a introdução começa em nova página e a capa é compactada para "Versão" ficar na 1ª página;
  - tabelas com largura de coluna proporcional ao texto (layout fixo), cabeçalho repetido em
    cada página, linhas que não se dividem e, até 9 linhas, tabela inteira junto da legenda.

Uso: python build_relatorio.py modelo.docx relatorio.md saida.docx capa.json [--repo] [--mapa ANTIGO=NOVO ...]

Caminhos das figuras:
  sem opções, o caminho escrito no Markdown é usado como está (absoluto, ou relativo à pasta do .md);
  --repo        resolve os caminhos absolutos do ambiente original (E:/CERISE/RELATORIOS/2026/... e
                E:/CERISE/FRENTES/...) para as cópias deste repositório (raiz e figuras/frentes/);
  --mapa A=B    troca o prefixo A por B (B relativo à pasta do .md, se não for absoluto); pode ser repetido.
"""
import argparse
import copy
import json
import os
import re
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageFont

AZUL = '1F4E79'
FONT = 'Arial'

# Prefixos do ambiente em que o relatório foi produzido -> pastas deste repositório (relativas à raiz do repositório)
RAIZ_REPO = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
MAPA_REPO = [('E:/CERISE/RELATORIOS/2026/', ''), ('E:/CERISE/FRENTES/', 'figuras/frentes/')]


def resolver_caminho(path, mapas, base_dir):
    """Aplica o primeiro prefixo de `mapas` (lista de pares antigo, novo) que casar com `path`.

    A comparação ignora maiúsculas/minúsculas e o tipo de barra; `novo` relativo é resolvido a partir de `base_dir`.
    Sem correspondência, devolve o caminho original (absoluto, ou relativo a `base_dir`)."""
    norm = path.replace('\\', '/')
    for antigo, novo in sorted(mapas, key=lambda par: -len(par[0])):
        a = antigo.replace('\\', '/')
        if norm.lower().startswith(a.lower()):
            resto = norm[len(a):].lstrip('/')
            destino = os.path.join(novo, resto) if novo else resto
            return os.path.normpath(destino if os.path.isabs(destino) else os.path.join(base_dir, destino))
    return path if os.path.isabs(path) else os.path.join(base_dir, path)


def set_run_font(run, size=None, bold=None, italic=None, color=None, mono=False):
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    f = 'Consolas' if mono else FONT
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(a), f)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, url, text, size=None):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), r_id)
    r = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), FONT)
    rpr.append(rf)
    c = OxmlElement('w:color'); c.set(qn('w:val'), '0563C1'); rpr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rpr.append(u)
    if size:
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2))); rpr.append(sz)
    r.append(rpr)
    t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve')
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|https?://[^\s)]+)')

# espaço entre número e unidade vira espaço não separável ("13,8 %" não quebra entre linhas); não vale para código
UNIDADE = re.compile(r'(?<=\d) (?=(?:%|°C|°|ms|s|min|h|B|KB|kB|KiB|MB|GB|TB|MiB|GiB|bytes|bits|bps|m|cm|mm|km|m²|m³|'
                     r'px|MP|ha|fps|FPS|Hz|kHz|MHz|GHz|dB|dBm|dBi|W|mW|V|mA|mAh|kg|g|rad|páginas|p\.p\.)(?![\wÀ-ÿ]))')
# outras quebras que separam um valor: "55,0 × 44,3", "ρ = 0,62", "8,3–8,6" e iniciais de nomes ("A. A. Cardoso")
NAO_QUEBRA = [(re.compile(r'(?<=\d) × (?=\d)'), '\u00a0× '),
              (re.compile(r'(?<=\S) = (?=[\d−-])'), '\u00a0=\u00a0'),
              (re.compile(r'(?<=\d)–(?=\d)'), '\u2060–\u2060'),
              (re.compile(r'(?<=\b[A-ZÀ-Ý]\.) (?=[A-ZÀ-Ý]\.)'), '\u00a0')]


def nbsp(texto):
    texto = UNIDADE.sub('\u00a0', texto)
    for padrao, troca in NAO_QUEBRA:
        texto = padrao.sub(troca, texto)
    return texto


_FONTES = {}


def _largura_twips(palavra, negrito=False, pt=9.5):
    """Largura de uma palavra em Arial (twips) no tamanho do texto das tabelas; 0 se a fonte não existir."""
    if negrito not in _FONTES:
        nome = 'arialbd.ttf' if negrito else 'arial.ttf'
        try:
            _FONTES[negrito] = ImageFont.truetype(os.path.join(os.environ.get('WINDIR', r'C:\Windows'), 'Fonts', nome), 1000)
        except OSError:
            _FONTES[negrito] = None
    f = _FONTES[negrito]
    return f.getlength(palavra) / 1000 * pt * 20 if f else 0


def add_inline(paragraph, text, size=None, base_bold=False, color=None):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            r = paragraph.add_run(nbsp(tok[2:-2])); set_run_font(r, size, bold=True, color=color)
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            r = paragraph.add_run(nbsp(tok[1:-1])); set_run_font(r, size, italic=True, bold=base_bold or None, color=color)
        elif tok.startswith('`') and tok.endswith('`'):
            r = paragraph.add_run(tok[1:-1]); set_run_font(r, (size or 12) - 1, mono=True, color=color)
        elif tok.startswith('[') and '](' in tok:
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', tok)
            add_hyperlink(paragraph, m.group(2), m.group(1), size)
        elif tok.startswith('http'):
            # pontuação final ("Disponível em: https://... . Acesso em") fica fora do link
            url = tok.rstrip('.,;:')
            add_hyperlink(paragraph, url, url, size)
            if url != tok:
                r = paragraph.add_run(tok[len(url):]); set_run_font(r, size, bold=base_bold or None, color=color)
        else:
            r = paragraph.add_run(nbsp(tok)); set_run_font(r, size, bold=base_bold or None, color=color)


def ppr_set(p, style=None, jc=None, first_line=None, left=None, hanging=None, before=None, after=None,
            line=276, keep_next=False, num=None, outline=None, page_break_before=False):
    pPr = p._p.get_or_add_pPr()
    if style:
        ps = OxmlElement('w:pStyle'); ps.set(qn('w:val'), style); pPr.insert(0, ps)
    if keep_next:  # ordem do esquema OOXML: keepNext, pageBreakBefore, widowControl, numPr
        pPr.append(OxmlElement('w:keepNext'))
    if page_break_before:
        pPr.append(OxmlElement('w:pageBreakBefore'))
    pPr.append(OxmlElement('w:widowControl'))  # sem linha sozinha no pé ou no topo da página
    if num:
        numPr = OxmlElement('w:numPr')
        il = OxmlElement('w:ilvl'); il.set(qn('w:val'), str(num[1])); numPr.append(il)
        ni = OxmlElement('w:numId'); ni.set(qn('w:val'), str(num[0])); numPr.append(ni)
        pPr.append(numPr)
    sp = OxmlElement('w:spacing')
    if before is not None:
        sp.set(qn('w:before'), str(before))
    sp.set(qn('w:after'), str(after if after is not None else 120))
    sp.set(qn('w:line'), str(line)); sp.set(qn('w:lineRule'), 'auto')
    pPr.append(sp)
    if first_line is not None or left is not None or hanging is not None:
        ind = OxmlElement('w:ind')
        if left is not None:
            ind.set(qn('w:left'), str(left))
        if first_line is not None:
            ind.set(qn('w:firstLine'), str(first_line))
        if hanging is not None:
            ind.set(qn('w:hanging'), str(hanging))
        pPr.append(ind)
    if jc:
        j = OxmlElement('w:jc'); j.set(qn('w:val'), jc); pPr.append(j)
    if outline is not None:
        o = OxmlElement('w:outlineLvl'); o.set(qn('w:val'), str(outline)); pPr.append(o)


class Builder:
    def __init__(self, template, md_path, capa, mapas=None):
        self.doc = docx.Document(template)
        self.md_dir = os.path.dirname(os.path.abspath(md_path))
        self.capa = capa
        self.mapas = mapas or []
        self.fig = 0
        self.tab = 0
        self.body = self.doc.element.body
        self.sect = self.body[-1]
        self._fix_heading_numbering()
        self._prepare()

    def _fix_heading_numbering(self):
        """Títulos numerados: número seguido de um espaço (sem tabulação) e sem recuo deslocado."""
        numbering = self.doc.part.numbering_part.element
        abs_id = None
        for num in numbering.findall(qn('w:num')):
            if num.get(qn('w:numId')) == '1':
                abs_id = num.find(qn('w:abstractNumId')).get(qn('w:val'))
        for an in numbering.findall(qn('w:abstractNum')):
            if an.get(qn('w:abstractNumId')) != abs_id:
                continue
            for lvl in an.findall(qn('w:lvl')):
                if lvl.get(qn('w:ilvl')) not in ('0', '1', '2'):
                    continue
                old = lvl.find(qn('w:suff'))
                if old is not None:
                    lvl.remove(old)
                suff = OxmlElement('w:suff'); suff.set(qn('w:val'), 'space')
                lvl.find(qn('w:lvlText')).addprevious(suff)
                ppr = lvl.find(qn('w:pPr'))
                ind = ppr.find(qn('w:ind')) if ppr is not None else None
                if ind is not None:
                    ind.set(qn('w:left'), '0'); ind.set(qn('w:hanging'), '0')

    # ---------- capa ----------
    def _set_para_text(self, el, label, value):
        p = docx.text.paragraph.Paragraph(el, self.doc)
        runs = p.runs
        fmt = copy.deepcopy(runs[0]._element.rPr) if runs and runs[0]._element.rPr is not None else None
        for r in runs:
            r._element.getparent().remove(r._element)
        for txt, bold in ((label, True), (value, True)):
            r = p.add_run(txt)
            if fmt is not None:
                r._element.insert(0, copy.deepcopy(fmt))
                hl = r._element.rPr.find(qn('w:highlight'))
                if hl is not None:
                    r._element.rPr.remove(hl)

    def _prepare(self):
        kids = list(self.body)
        texts = [''.join(t.text or '' for t in k.iter() if t.tag == qn('w:t')) for k in kids]
        idx_intro = next(i for i, t in enumerate(texts) if t.strip().startswith('INTRODUÇÃO'))
        for i, t in enumerate(texts[:idx_intro]):
            if t.startswith('Título do Plano de Trabalho Individual'):
                self._set_para_text(kids[i], 'Título do Plano de Trabalho Individual: ', self.capa['titulo'])
            elif t.startswith('Grupo / Eixo Temático'):
                self._set_para_text(kids[i], 'Grupo / Eixo Temático: ', self.capa['grupo'])
            elif t.startswith('Coordenador do Grupo'):
                self._set_para_text(kids[i], self.capa.get('rotulo_coord', 'Coordenadores dos Grupos / Eixos Temáticos: '), self.capa['coordenador'])
            elif t.startswith('Versão:'):
                self._set_para_text(kids[i], 'Versão: ', self.capa['versao'])
            elif t.strip() == 'Relatório Técnico-Científico' and self.capa.get('tipo'):
                # permite reaproveitar o modelo para outros documentos (ex.: "Roteiro de Implementação")
                self._set_para_text(kids[i], '', self.capa['tipo'])
        if self.capa.get('bolsista'):
            # insere o nome do bolsista logo após o título do plano
            anchor = next(k for k, t in zip(kids, texts) if t.startswith('Título do Plano'))
            p = OxmlElement('w:p'); anchor.addnext(p)
            par = docx.text.paragraph.Paragraph(p, self.doc)
            ppr_set(par, jc='center', before=240, after=0)
            r = par.add_run('Bolsista: ' + self.capa['bolsista']); set_run_font(r, 12, bold=True, color='404040')
        # Os textos da capa (título, coordenadores, bolsista) ocupam mais linhas que os "XXXXX" do modelo:
        # remove parágrafos vazios entre os blocos (mantendo um como respiro) para "Versão" ficar na 1ª página.
        def vazio(i):
            k = kids[i]
            return (not texts[i].strip() and k.tag == qn('w:p')
                    and not any(e.tag in (qn('w:drawing'), qn('w:br')) for e in k.iter()))
        marcos = [i for i, t in enumerate(texts[:idx_intro])
                  if t.startswith(('Título do Plano', 'Grupo / Eixo', 'Coordenador do Grupo', 'Versão:'))]
        for a, b in zip(marcos, marcos[1:]):
            for i in [i for i in range(a + 1, b) if vazio(i)][1:]:
                self.body.remove(kids[i])
        idx_versao = marcos[-1]
        for i in range(idx_versao + 1, idx_intro):  # vazios após "Versão": a introdução começa em nova página
            if vazio(i):
                self.body.remove(kids[i])
        for k in kids[idx_intro:-1]:
            self.body.remove(k)
        self.first_h1 = True

    # ---------- blocos ----------
    def _new_p(self):
        p = OxmlElement('w:p')
        self.sect.addprevious(p)
        return docx.text.paragraph.Paragraph(p, self.doc)

    @staticmethod
    def _mark_rpr(p, size, color):
        """Formata a marca de parágrafo (é ela que define a fonte do número automático do título)."""
        pPr = p._p.get_or_add_pPr()
        rpr = OxmlElement('w:rPr')
        rf = OxmlElement('w:rFonts')
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(a), FONT)
        rpr.append(rf)
        rpr.append(OxmlElement('w:b'))
        c = OxmlElement('w:color'); c.set(qn('w:val'), color); rpr.append(c)
        for tag in ('w:sz', 'w:szCs'):
            e = OxmlElement(tag); e.set(qn('w:val'), str(int(size * 2))); rpr.append(e)
        pPr.append(rpr)

    def heading(self, level, text):
        p = self._new_p()
        if level == 1:
            ppr_set(p, style='Ttulo1', num=(1, 0), before=360, after=160, jc='left', keep_next=True,
                    page_break_before=getattr(self, 'first_h1', False))
            self.first_h1 = False
            self._mark_rpr(p, 14, AZUL)
            add_inline(p, text.upper(), 14, base_bold=True, color=AZUL)
        else:
            color = AZUL if level == 2 else '2E2E2E'
            ppr_set(p, style='Ttulo2', num=(1, level - 1), before=240, after=120, jc='left', keep_next=True,
                    outline=level - 1)
            self._mark_rpr(p, 12, color)
            add_inline(p, text, 12, base_bold=True, color=color)

    def para(self, text):
        p = self._new_p()
        ppr_set(p, style='Normal1', first_line=709, jc='both', after=120)
        add_inline(p, text, 12)

    def ref(self, text):
        p = self._new_p()
        ppr_set(p, style='Normal1', left=567, hanging=567, jc='left', after=60, line=240)
        add_inline(p, text, 11)

    def bullet(self, text, level=0):
        p = self._new_p()
        ppr_set(p, style='Normal1', num=(2, level), jc='both', after=60, left=1077 + 360 * level, hanging=357)
        add_inline(p, text, 12)

    def code(self, linhas):
        """Bloco de código (```): fonte monoespaçada (8 pt, ~93 caracteres por linha), fundo cinza claro.

        Linha que não cabe na largura continua abaixo com recuo deslocado (quebra só visual: copiada do PDF,
        volta a ser uma linha só)."""
        for k, ln in enumerate(linhas):
            p = self._new_p()
            ppr_set(p, jc='left', left=284 + 454, hanging=454, after=0 if k < len(linhas) - 1 else 140, line=240,
                    keep_next=k < len(linhas) - 1)
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), 'F2F2F2')
            pPr.find(qn('w:spacing')).addprevious(shd)  # ordem do esquema OOXML: shd antes de spacing
            r = p.add_run(ln if ln else ' ')
            r._element.find(qn('w:t')).set(qn('xml:space'), 'preserve')
            set_run_font(r, 8, mono=True, color='1F1F1F')

    def pagebreak(self):
        p = self._new_p()
        p.add_run().add_break(WD_BREAK.PAGE)

    def figure(self, path, caption, width_cm=15.0):
        full = resolver_caminho(path, self.mapas, self.md_dir)
        if not os.path.exists(full):
            self.para(f'[FIGURA AUSENTE: {path}]')
            print('AVISO: figura ausente', full, file=sys.stderr)
            return
        with Image.open(full) as im:
            w, h = im.size
        width = min(width_cm, 16.0)
        if h / w * width > 19.0:  # limita a altura
            width = 19.0 * w / h
        p = self._new_p()
        ppr_set(p, jc='center', before=120, after=60, keep_next=True)
        p.add_run().add_picture(full, width=Cm(width))
        self.fig += 1
        c = self._new_p()
        ppr_set(c, jc='center', after=200)
        r = c.add_run(f'Figura {self.fig}. '); set_run_font(r, 10, bold=True)
        add_inline(c, caption, 10)

    def table(self, rows, caption=None):
        if caption:
            self.tab += 1
            c = self._new_p()
            ppr_set(c, jc='center', before=160, after=60, keep_next=True)
            r = c.add_run(f'Tabela {self.tab}. '); set_run_font(r, 10, bold=True)
            add_inline(c, caption, 10)
        ncol = max(len(r) for r in rows)
        widths = self._col_widths(rows, ncol)
        # tabelas curtas ficam inteiras numa página, junto com a legenda; nas longas, a legenda fica ao menos com
        # o cabeçalho e a primeira linha de dados (o cabeçalho se repete nas páginas seguintes)
        keep_rows = len(rows) - 1 if len(rows) <= 9 else 2
        t = self.doc.add_table(rows=len(rows), cols=ncol)
        self.sect.addprevious(t._tbl)
        tblPr = t._tbl.tblPr
        tw = tblPr.find(qn('w:tblW'))
        if tw is None:
            tw = OxmlElement('w:tblW'); tblPr.append(tw)
        tw.set(qn('w:w'), str(sum(widths))); tw.set(qn('w:type'), 'dxa')
        borders = OxmlElement('w:tblBorders')
        for b in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            e = OxmlElement(f'w:{b}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), '808080')
            borders.append(e)
        tblPr.append(borders)
        lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
        jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); tblPr.append(jc)
        for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
            gc.set(qn('w:w'), str(w))
        for i, row in enumerate(rows):
            trPr = t.rows[i]._tr.get_or_add_trPr()
            trPr.append(OxmlElement('w:cantSplit'))
            if i == 0:
                trPr.append(OxmlElement('w:tblHeader'))
            for j in range(ncol):
                cell = t.cell(i, j)
                txt = row[j] if j < len(row) else ''
                p = cell.paragraphs[0]
                ppr_set(p, jc='center' if i == 0 else 'left', after=0, line=240,
                        keep_next=i < keep_rows)
                add_inline(p, txt.strip(), 9.5, base_bold=(i == 0), color=None)
                tcPr = cell._tc.get_or_add_tcPr()
                tcw = tcPr.find(qn('w:tcW'))
                if tcw is None:
                    tcw = OxmlElement('w:tcW'); tcPr.insert(0, tcw)
                tcw.set(qn('w:w'), str(widths[j])); tcw.set(qn('w:type'), 'dxa')
                if i == 0:
                    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), 'DCE6F1')
                    tcPr.append(shd)
        sp = self._new_p(); ppr_set(sp, after=120)

    @staticmethod
    def _plain(txt):
        txt = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', txt)
        return re.sub(r'[*`]', '', txt).strip()

    def _col_widths(self, rows, ncol, total=8504):
        """Largura de cada coluna (twips) proporcional ao texto, com mínimo de 1,5 cm e da maior palavra."""
        CH = 110            # largura média de um caractere em Arial 9,5 pt (twips), com folga para negrito
        PAD = 230           # margens internas da célula
        desired, minimum = [], []
        for j in range(ncol):
            cells = [self._plain(r[j]) if j < len(r) else '' for r in rows]
            body = cells[1:] or cells
            longest = max(len(c) for c in body)
            head = len(cells[0])
            word = max((len(w) for c in cells for w in c.split()), default=1)
            # opcional (capa.json "tabela_max_palavra"): limita o peso de palavras muito longas, como caminhos de
            # arquivo, que então quebram dentro da célula em vez de estreitar as outras colunas
            if self.capa.get('tabela_max_palavra'):
                word = min(word, int(self.capa['tabela_max_palavra']))
            desired.append(max(min(longest, 100), 0.55 * head, word) * CH + PAD)
            # maior palavra medida na fonte real (cabeçalho em negrito), para não sair "Makespa|n" nem "ADD@6|0";
            # 216 twips de margens da célula e 40 de folga
            # trechos unidos por espaço não separável ("60 mm") contam como uma palavra só
            medida = max([_largura_twips(w, True) for w in nbsp(cells[0]).split(' ')] +
                         [_largura_twips(w) for c in body for w in nbsp(c).split(' ')] + [0]) + 256
            if self.capa.get('tabela_max_palavra'):
                medida = min(medida, word * CH + PAD)
            # células curtas (até 12 caracteres) não quebram; nas demais, basta caber a maior palavra
            minimum.append(max(850, word * CH + PAD, medida, (longest * CH + PAD) if longest <= 12 else 0))
        if sum(minimum) >= total:
            return [int(m * total / sum(minimum)) for m in minimum]
        scale = total / sum(desired)
        w = [max(m, d * scale) for m, d in zip(minimum, desired)]
        for _ in range(10):  # devolve o excesso proporcionalmente às colunas acima do mínimo
            excess = sum(w) - total
            if excess <= 1:
                break
            flex = [x - m for x, m in zip(w, minimum)]
            tf = sum(flex)
            if tf <= 0:
                break
            w = [x - excess * f / tf for x, f in zip(w, flex)]
        return [int(x) for x in w]

    # ---------- markdown ----------
    def run_md(self, md):
        lines = md.splitlines()
        i, buf, pending_caption = 0, [], None

        def flush():
            if buf:
                text = ' '.join(s.strip() for s in buf)
                if re.match(r'^\[\d+\]', text):
                    self.ref(text)
                else:
                    self.para(text)
                buf.clear()

        while i < len(lines):
            ln = lines[i].rstrip()
            s = ln.strip()
            if not s:
                flush(); i += 1; continue
            if s.startswith('```'):
                flush()
                bloco = []
                i += 1
                while i < len(lines) and not lines[i].strip().startswith('```'):
                    bloco.append(lines[i].rstrip('\n').expandtabs(4))
                    i += 1
                i += 1
                self.code(bloco)
                continue
            if s == '<!-- pagebreak -->':
                flush(); self.pagebreak(); i += 1; continue
            if s.startswith('<!--'):
                flush()
                while i < len(lines) and '-->' not in lines[i]:
                    i += 1
                i += 1; continue
            m = re.match(r'^(#{1,3})\s+(.*)$', s)
            if m:
                flush(); self.heading(len(m.group(1)), m.group(2).strip()); i += 1; continue
            m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)(\{w=([\d.]+)\})?\s*$', s)
            if m:
                flush(); self.figure(m.group(2), m.group(1), float(m.group(4)) if m.group(4) else 15.0); i += 1; continue
            m = re.match(r'^Tabela:\s*(.*)$', s)
            if m and i + 1 < len(lines) and lines[i + 1].strip().startswith('|'):
                flush(); pending_caption = m.group(1); i += 1; continue
            if s.startswith('|'):
                flush()
                rows = []
                while i < len(lines) and lines[i].strip().startswith('|'):
                    cells = [c for c in lines[i].strip().strip('|').split('|')]
                    if not all(re.match(r'^\s*:?-{2,}:?\s*$', c) for c in cells):
                        rows.append(cells)
                    i += 1
                self.table(rows, pending_caption); pending_caption = None
                continue
            m = re.match(r'^(\s*)[-*]\s+(.*)$', ln)
            if m:
                flush(); self.bullet(m.group(2), min(len(m.group(1)) // 2, 2)); i += 1; continue
            buf.append(s); i += 1
        flush()

    def save(self, out):
        import datetime
        cp = self.doc.core_properties
        cp.title = self.capa.get('titulo_documento', 'Relatório Técnico-Científico CERISE')
        cp.author = self.capa.get('bolsista', '')
        # o modelo oficial traz metadados do arquivo de origem (outro autor, datas de 2024); substitui pelos do relatório
        cp.last_modified_by = self.capa.get('bolsista', '')
        agora = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0)
        cp.created = agora
        cp.modified = agora
        cp.revision = 1
        self.doc.save(out)


def main():
    ap = argparse.ArgumentParser(description='Gera o relatório CERISE (.docx) a partir de um Markdown.')
    ap.add_argument('template'); ap.add_argument('md_path'); ap.add_argument('out'); ap.add_argument('capa_path')
    ap.add_argument('--repo', action='store_true',
                    help='resolve os caminhos absolutos E:/CERISE/... das figuras para as cópias deste repositório')
    ap.add_argument('--mapa', action='append', default=[], metavar='ANTIGO=NOVO',
                    help='troca o prefixo ANTIGO por NOVO nos caminhos das figuras (pode ser repetido)')
    a = ap.parse_args()
    mapas = [tuple(m.split('=', 1)) for m in a.mapa]
    if a.repo:
        mapas += [(antigo, os.path.join(RAIZ_REPO, novo)) for antigo, novo in MAPA_REPO]
    template, md_path, out, capa_path = a.template, a.md_path, a.out, a.capa_path
    capa = json.load(open(capa_path, encoding='utf-8'))
    b = Builder(template, md_path, capa, mapas)
    b.run_md(open(md_path, encoding='utf-8').read())
    b.save(out)
    print(f'OK {out} | figuras={b.fig} tabelas={b.tab}')


if __name__ == '__main__':
    main()
