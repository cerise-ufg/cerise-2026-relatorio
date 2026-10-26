#!/usr/bin/env python3
"""Gera o relatório CERISE (.docx) a partir de um Markdown, usando o modelo oficial como base.

Mantém a capa, o cabeçalho com os logos, a numeração automática dos títulos (1., 1.1., 1.1.1.)
e o estilo de parágrafo do modelo (Arial 12, justificado, recuo de 1,25 cm na primeira linha).

Markdown suportado:
  # / ## / ###            títulos numerados (níveis 1-3)
  parágrafos              linhas consecutivas viram um parágrafo; linha em branco separa
  - item / * item         marcadores
  ![legenda](caminho){w=14}   figura centralizada (largura em cm, padrão 15) + "Figura N. legenda"
  Tabela: legenda         linha logo antes de uma tabela pipe -> "Tabela N. legenda" acima dela
  | a | b |               tabela pipe (1ª linha = cabeçalho; linha |---| ignorada)
  **negrito** *itálico* `código` [texto](url)
  <!-- pagebreak -->      quebra de página
  [1] referência          parágrafo de referência (sem recuo)

Uso: python build_relatorio.py modelo.docx relatorio.md saida.docx capa.json
"""
import copy
import json
import os
import re
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

AZUL = '1F4E79'
FONT = 'Arial'


def set_run_font(run, size=None, bold=None, italic=None, color=None, mono=False):
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    f = 'Consolas' if mono else FONT
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(a), f)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, url, text, size=None):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), r_id)
    r = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), FONT)
    rpr.append(rf)
    c = OxmlElement('w:color'); c.set(qn('w:val'), '0563C1'); rpr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rpr.append(u)
    if size:
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2))); rpr.append(sz)
    r.append(rpr)
    t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve')
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|https?://[^\s)]+)')


def add_inline(paragraph, text, size=None, base_bold=False, color=None):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            r = paragraph.add_run(tok[2:-2]); set_run_font(r, size, bold=True, color=color)
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            r = paragraph.add_run(tok[1:-1]); set_run_font(r, size, italic=True, bold=base_bold or None, color=color)
        elif tok.startswith('`') and tok.endswith('`'):
            r = paragraph.add_run(tok[1:-1]); set_run_font(r, (size or 12) - 1, mono=True, color=color)
        elif tok.startswith('[') and '](' in tok:
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', tok)
            add_hyperlink(paragraph, m.group(2), m.group(1), size)
        elif tok.startswith('http'):
            add_hyperlink(paragraph, tok, tok, size)
        else:
            r = paragraph.add_run(tok); set_run_font(r, size, bold=base_bold or None, color=color)


def ppr_set(p, style=None, jc=None, first_line=None, left=None, hanging=None, before=None, after=None,
            line=276, keep_next=False, num=None, outline=None):
    pPr = p._p.get_or_add_pPr()
    if style:
        ps = OxmlElement('w:pStyle'); ps.set(qn('w:val'), style); pPr.insert(0, ps)
    if num:
        numPr = OxmlElement('w:numPr')
        il = OxmlElement('w:ilvl'); il.set(qn('w:val'), str(num[1])); numPr.append(il)
        ni = OxmlElement('w:numId'); ni.set(qn('w:val'), str(num[0])); numPr.append(ni)
        pPr.append(numPr)
    if keep_next:
        pPr.append(OxmlElement('w:keepNext'))
    sp = OxmlElement('w:spacing')
    if before is not None:
        sp.set(qn('w:before'), str(before))
    sp.set(qn('w:after'), str(after if after is not None else 120))
    sp.set(qn('w:line'), str(line)); sp.set(qn('w:lineRule'), 'auto')
    pPr.append(sp)
    if first_line is not None or left is not None or hanging is not None:
        ind = OxmlElement('w:ind')
        if left is not None:
            ind.set(qn('w:left'), str(left))
        if first_line is not None:
            ind.set(qn('w:firstLine'), str(first_line))
        if hanging is not None:
            ind.set(qn('w:hanging'), str(hanging))
        pPr.append(ind)
    if jc:
        j = OxmlElement('w:jc'); j.set(qn('w:val'), jc); pPr.append(j)
    if outline is not None:
        o = OxmlElement('w:outlineLvl'); o.set(qn('w:val'), str(outline)); pPr.append(o)


class Builder:
    def __init__(self, template, md_path, capa):
        self.doc = docx.Document(template)
        self.md_dir = os.path.dirname(os.path.abspath(md_path))
        self.capa = capa
        self.fig = 0
        self.tab = 0
        self.body = self.doc.element.body
        self.sect = self.body[-1]
        self._prepare()

    # ---------- capa ----------
    def _set_para_text(self, el, label, value):
        p = docx.text.paragraph.Paragraph(el, self.doc)
        runs = p.runs
        fmt = copy.deepcopy(runs[0]._element.rPr) if runs and runs[0]._element.rPr is not None else None
        for r in runs:
            r._element.getparent().remove(r._element)
        for txt, bold in ((label, True), (value, True)):
            r = p.add_run(txt)
            if fmt is not None:
                r._element.insert(0, copy.deepcopy(fmt))
                hl = r._element.rPr.find(qn('w:highlight'))
                if hl is not None:
                    r._element.rPr.remove(hl)

    def _prepare(self):
        kids = list(self.body)
        texts = [''.join(t.text or '' for t in k.iter() if t.tag == qn('w:t')) for k in kids]
        idx_intro = next(i for i, t in enumerate(texts) if t.strip().startswith('INTRODUÇÃO'))
        for i, t in enumerate(texts[:idx_intro]):
            if t.startswith('Título do Plano de Trabalho Individual'):
                self._set_para_text(kids[i], 'Título do Plano de Trabalho Individual: ', self.capa['titulo'])
            elif t.startswith('Grupo / Eixo Temático'):
                self._set_para_text(kids[i], 'Grupo / Eixo Temático: ', self.capa['grupo'])
            elif t.startswith('Coordenador do Grupo'):
                self._set_para_text(kids[i], self.capa.get('rotulo_coord', 'Coordenadores dos Grupos / Eixos Temáticos: '), self.capa['coordenador'])
            elif t.startswith('Versão:'):
                self._set_para_text(kids[i], 'Versão: ', self.capa['versao'])
        if self.capa.get('bolsista'):
            # insere o nome do bolsista logo após o título do plano
            anchor = next(k for k, t in zip(kids, texts) if t.startswith('Título do Plano'))
            p = OxmlElement('w:p'); anchor.addnext(p)
            par = docx.text.paragraph.Paragraph(p, self.doc)
            ppr_set(par, jc='center', before=240, after=0)
            r = par.add_run('Bolsista: ' + self.capa['bolsista']); set_run_font(r, 12, bold=True, color='404040')
            # compensa a linha nova removendo parágrafos vazios da capa, para "Versão" continuar na 1ª página
            vazios = [k for k, t in zip(kids[:idx_intro], texts[:idx_intro])
                      if not t.strip() and k.tag == qn('w:p') and not any(e.tag == qn('w:drawing') for e in k.iter())]
            for k in vazios[-6:-1]:
                self.body.remove(k)
        for k in kids[idx_intro:-1]:
            self.body.remove(k)

    # ---------- blocos ----------
    def _new_p(self):
        p = OxmlElement('w:p')
        self.sect.addprevious(p)
        return docx.text.paragraph.Paragraph(p, self.doc)

    @staticmethod
    def _mark_rpr(p, size, color):
        """Formata a marca de parágrafo (é ela que define a fonte do número automático do título)."""
        pPr = p._p.get_or_add_pPr()
        rpr = OxmlElement('w:rPr')
        rf = OxmlElement('w:rFonts')
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(a), FONT)
        rpr.append(rf)
        rpr.append(OxmlElement('w:b'))
        c = OxmlElement('w:color'); c.set(qn('w:val'), color); rpr.append(c)
        for tag in ('w:sz', 'w:szCs'):
            e = OxmlElement(tag); e.set(qn('w:val'), str(int(size * 2))); rpr.append(e)
        pPr.append(rpr)

    def heading(self, level, text):
        p = self._new_p()
        if level == 1:
            ppr_set(p, style='Ttulo1', num=(1, 0), before=360, after=160, jc='both', keep_next=True)
            self._mark_rpr(p, 14, AZUL)
            add_inline(p, text.upper(), 14, base_bold=True, color=AZUL)
        else:
            color = AZUL if level == 2 else '2E2E2E'
            ppr_set(p, style='Ttulo2', num=(1, level - 1), before=240, after=120, jc='both', keep_next=True,
                    outline=level - 1)
            self._mark_rpr(p, 12, color)
            add_inline(p, text, 12, base_bold=True, color=color)

    def para(self, text):
        p = self._new_p()
        ppr_set(p, style='Normal1', first_line=709, jc='both', after=120)
        add_inline(p, text, 12)

    def ref(self, text):
        p = self._new_p()
        ppr_set(p, style='Normal1', left=567, hanging=567, jc='both', after=100)
        add_inline(p, text, 11)

    def bullet(self, text, level=0):
        p = self._new_p()
        ppr_set(p, style='Normal1', num=(2, level), jc='both', after=60, left=1077 + 360 * level, hanging=357)
        add_inline(p, text, 12)

    def pagebreak(self):
        p = self._new_p()
        p.add_run().add_break(WD_BREAK.PAGE)

    def figure(self, path, caption, width_cm=15.0):
        full = path if os.path.isabs(path) else os.path.join(self.md_dir, path)
        if not os.path.exists(full):
            self.para(f'[FIGURA AUSENTE: {path}]')
            print('AVISO: figura ausente', full, file=sys.stderr)
            return
        with Image.open(full) as im:
            w, h = im.size
        width = min(width_cm, 16.0)
        if h / w * width > 19.0:  # limita a altura
            width = 19.0 * w / h
        p = self._new_p()
        ppr_set(p, jc='center', before=120, after=60, keep_next=True)
        p.add_run().add_picture(full, width=Cm(width))
        self.fig += 1
        c = self._new_p()
        ppr_set(c, jc='center', after=200)
        r = c.add_run(f'Figura {self.fig}. '); set_run_font(r, 10, bold=True)
        add_inline(c, caption, 10)

    def table(self, rows, caption=None):
        if caption:
            self.tab += 1
            c = self._new_p()
            ppr_set(c, jc='center', before=160, after=60, keep_next=True)
            r = c.add_run(f'Tabela {self.tab}. '); set_run_font(r, 10, bold=True)
            add_inline(c, caption, 10)
        ncol = max(len(r) for r in rows)
        t = self.doc.add_table(rows=len(rows), cols=ncol)
        self.sect.addprevious(t._tbl)
        tblPr = t._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for b in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            e = OxmlElement(f'w:{b}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), '808080')
            borders.append(e)
        tblPr.append(borders)
        jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); tblPr.append(jc)
        for i, row in enumerate(rows):
            for j in range(ncol):
                cell = t.cell(i, j)
                txt = row[j] if j < len(row) else ''
                p = cell.paragraphs[0]
                ppr_set(p, jc='center' if i == 0 else 'left', after=0, line=240)
                add_inline(p, txt.strip(), 9.5, base_bold=(i == 0), color=None)
                if i == 0:
                    tcPr = cell._tc.get_or_add_tcPr()
                    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), 'DCE6F1')
                    tcPr.append(shd)
        sp = self._new_p(); ppr_set(sp, after=120)

    # ---------- markdown ----------
    def run_md(self, md):
        lines = md.splitlines()
        i, buf, pending_caption = 0, [], None

        def flush():
            if buf:
                text = ' '.join(s.strip() for s in buf)
                if re.match(r'^\[\d+\]', text):
                    self.ref(text)
                else:
                    self.para(text)
                buf.clear()

        while i < len(lines):
            ln = lines[i].rstrip()
            s = ln.strip()
            if not s:
                flush(); i += 1; continue
            if s == '<!-- pagebreak -->':
                flush(); self.pagebreak(); i += 1; continue
            if s.startswith('<!--'):
                flush()
                while i < len(lines) and '-->' not in lines[i]:
                    i += 1
                i += 1; continue
            m = re.match(r'^(#{1,3})\s+(.*)$', s)
            if m:
                flush(); self.heading(len(m.group(1)), m.group(2).strip()); i += 1; continue
            m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)(\{w=([\d.]+)\})?\s*$', s)
            if m:
                flush(); self.figure(m.group(2), m.group(1), float(m.group(4)) if m.group(4) else 15.0); i += 1; continue
            m = re.match(r'^Tabela:\s*(.*)$', s)
            if m and i + 1 < len(lines) and lines[i + 1].strip().startswith('|'):
                flush(); pending_caption = m.group(1); i += 1; continue
            if s.startswith('|'):
                flush()
                rows = []
                while i < len(lines) and lines[i].strip().startswith('|'):
                    cells = [c for c in lines[i].strip().strip('|').split('|')]
                    if not all(re.match(r'^\s*:?-{2,}:?\s*$', c) for c in cells):
                        rows.append(cells)
                    i += 1
                self.table(rows, pending_caption); pending_caption = None
                continue
            m = re.match(r'^(\s*)[-*]\s+(.*)$', ln)
            if m:
                flush(); self.bullet(m.group(2), min(len(m.group(1)) // 2, 2)); i += 1; continue
            buf.append(s); i += 1
        flush()

    def save(self, out):
        self.doc.core_properties.title = 'Relatório Técnico-Científico CERISE'
        self.doc.core_properties.author = self.capa.get('bolsista', '')
        self.doc.save(out)


def main():
    template, md_path, out, capa_path = sys.argv[1:5]
    capa = json.load(open(capa_path, encoding='utf-8'))
    b = Builder(template, md_path, capa)
    b.run_md(open(md_path, encoding='utf-8').read())
    b.save(out)
    print(f'OK {out} | figuras={b.fig} tabelas={b.tab}')


if __name__ == '__main__':
    main()
